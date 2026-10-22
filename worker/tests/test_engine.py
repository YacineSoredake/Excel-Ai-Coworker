"""Benchmark (report §22): representative workbooks with expected interpretations."""
import os
from datetime import date, datetime

import openpyxl
import pytest
from openpyxl.styles import Font
from openpyxl.worksheet.table import Table

from app.engine.analyze import analyze_workbook
from app.engine.context import Limits
from app.engine.formulas import inconsistent_formulas, relativize
from app.validate import Rejected

from .helpers import build, codes, col, finding

FIX = os.path.join(os.path.dirname(__file__), "fixtures")


def load(name):
    return analyze_workbook(open(os.path.join(FIX, name), "rb").read())


# ------------------------------------------------------------------ the two files the user tested with

def test_clean_sales_interpretation():
    out = load("01_sales_clean.xlsx")
    assert [t.name for t in out.tables] == ["Sales", "Targets", "Lookup"]
    sales = out.tables[0]
    assert (sales.header_row, sales.first_data_row, len(sales.rows)) == (1, 2, 339)
    c = col(out)
    exp = {
        "order_id": ("identifier", "identifier", None), "date": ("date", "time", "date"),
        "wilaya": ("text", "dimension", "location"), "category": ("text", "dimension", "category"),
        "product": ("text", "dimension", "product"), "quantity": ("integer", "measure", "quantity"),
        "unit_price_dzd": ("currency", "measure", "price"), "total_dzd": ("currency", "measure", "amount"),
        "channel": ("text", "dimension", "category"), "payment_method": ("text", "dimension", "category"),
    }
    for name, (dtype, role, sub) in exp.items():
        assert (c[name].data_type, c[name].role, c[name].role_subtype) == (dtype, role, sub), name
    assert c["total_dzd"].currency == "DZD"
    assert c["date"].stats["granularity"] in ("daily", "weekly")


def test_clean_sales_is_quiet_and_finds_the_lookup_relationship():
    out = load("01_sales_clean.xlsx")
    assert {f.severity for f in out.findings} <= {"info"}
    assert codes(out) == ["hidden_sheet"]
    rel = out.relationships[0]
    assert (rel["from_table_name"], rel["from_column"], rel["to_table_name"]) == ("Sales", "wilaya", "Lookup")


def test_messy_french_interpretation():
    out = load("02_messy_ventes_fr.xlsx")
    t = out.tables[0]
    assert (t.header_row, t.first_data_row, t.total_rows) == (4, 5, [66])
    assert len(t.rows) == 61
    assert out.sheets[0]["notes"][0].startswith("Rapport des ventes")
    c = col(out)
    assert c["Date"].data_type == "date"
    assert c["Qté"].data_type == "integer"
    assert c["Prix unitaire (DA)"].currency == "DZD" and c["Montant (DA)"].currency == "DZD"
    assert (c["Statut"].role, c["Statut"].role_subtype) == ("dimension", "status")
    assert c["Client"].role_subtype == "person"


def test_messy_french_findings():
    out = load("02_messy_ventes_fr.xlsx")
    assert finding(out, "duplicate_rows").sample_rows == [65]
    assert finding(out, "missing_values", "Client").affected == 7
    assert finding(out, "missing_values", "Qté").affected == 6
    assert finding(out, "inconsistent_labels", "Wilaya") is not None
    assert finding(out, "mixed_date_formats", "Date") is not None
    o = finding(out, "extreme_outliers", "Montant (DA)")
    assert o and o.sample_rows == [25] and o.evidence["examples"][0]["value"] == 9999999.0
    assert finding(out, "suspected_total_row").sample_rows == [66]
    assert finding(out, "mostly_empty_column", "Remarque") is not None
    assert all("Potential issue" in f.title for f in out.findings if f.severity in ("medium", "high"))


def test_blank_row_inside_table_is_one_table():
    out = load("03_blank_row_inside_table.xlsx")
    assert len(out.tables) == 1 and out.tables[0].blank_rows == 1
    assert len(out.tables[0].rows) == 60
    assert "blank_rows_inside" in codes(out)


# ------------------------------------------------------------------ adversarial workbooks

def test_two_tables_on_one_sheet():
    def f(wb):
        ws = wb.active
        ws.append(["Produit", "Ventes"])
        for i in range(12):
            ws.append([f"P{i}", 100 + i * 7])
        ws.append([])
        ws.append([])
        ws.append(["Région", "Objectif"])
        for i in range(8):
            ws.append([f"R{i}", 1000 + i])
    out = analyze_workbook(build(f))
    assert len(out.tables) == 2
    assert out.tables[0].name.endswith("table 1") and out.tables[1].header_row == 16
    assert len(out.tables[0].rows) == 12 and len(out.tables[1].rows) == 8


def test_title_directly_above_header_without_blank_row():
    def f(wb):
        ws = wb.active
        ws["A1"] = "Stock au 30 septembre"
        ws.append(["Article", "Quantité", "Prix", "Fournisseur"])
        for i in range(15):
            ws.append([f"Art {i}", i + 1, 100.5 + i, f"F{i % 3}"])
    out = analyze_workbook(build(f))
    t = out.tables[0]
    assert t.header_row == 2 and len(t.rows) == 15
    assert "Stock au 30 septembre" in t.preamble


def test_numbers_stored_as_text_french_formats():
    def f(wb):
        ws = wb.active
        ws.append(["Client", "Montant", "Remise", "Téléphone"])
        for i in range(12):
            ws.append([f"C{i}", f"{1000 + i * 250:,}".replace(",", " ") + ",50 DA", f"{i + 1} %", f"055512{i:04d}"])
    out = analyze_workbook(build(f))
    c = col(out)
    assert c["Montant"].data_type == "currency" and c["Montant"].currency == "DZD"
    assert c["Montant"].values[0] == pytest.approx(1000.5)
    assert c["Remise"].data_type == "percentage" and c["Remise"].values[0] == pytest.approx(0.01)
    assert c["Téléphone"].base_type == "text"  # leading zeros kept as text, not turned into numbers


def test_excel_defined_table_is_trusted():
    def f(wb):
        ws = wb.active
        ws["C3"], ws["D3"] = "Nom", "Score"
        for i in range(10):
            ws.cell(row=4 + i, column=3, value=f"N{i}")
            ws.cell(row=4 + i, column=4, value=i * 3)
        t = Table(displayName="Scores", ref="C3:D13")
        ws.add_table(t)
    out = analyze_workbook(build(f))
    t = out.tables[0]
    assert t.source == "excel_table" and t.name == "Scores" and t.confidence == 0.99
    assert t.cell_range == "C3:D13"


def test_conflicting_date_order_is_flagged_high():
    def f(wb):
        ws = wb.active
        ws.append(["Date", "Montant"])
        for d in ["25/12/2026", "01/13/2026", "03/04/2026", "05/06/2026", "28/02/2026", "02/28/2026"]:
            ws.append([d, 100])
    out = analyze_workbook(build(f))
    fd = finding(out, "mixed_date_formats", "Date")
    assert fd and fd.severity == "high"


def test_all_month_first_dates_are_read_month_first():
    def f(wb):
        ws = wb.active
        ws.append(["Date", "Montant"])
        for d in ["12/25/2026", "03/04/2026", "01/31/2026", "05/06/2026", "07/08/2026"]:
            ws.append([d, 100])
    out = analyze_workbook(build(f))
    c = col(out)["Date"]
    assert c.date_info["order_assumed"] == "mdy"
    assert c.values[1].month == 3 and c.values[1].day == 4


def test_arabic_headers_and_values():
    def f(wb):
        ws = wb.active
        ws.append(["المنتج", "الكمية", "السعر", "ولاية"])
        for i in range(12):
            ws.append([f"منتج {i % 3}", i + 1, 100 + i, ["الجزائر", "وهران", "قسنطينة"][i % 3]])
    out = analyze_workbook(build(f))
    c = col(out)
    assert c["الكمية"].role_subtype == "quantity" and c["السعر"].role_subtype == "price"
    assert c["ولاية"].role_subtype == "location" and c["المنتج"].role_subtype == "product"


def test_empty_header_only_and_notes_sheets():
    def f(wb):
        wb.active.title = "Vide"
        ws = wb.create_sheet("Entête")
        ws.append(["a", "b", "c"])
        ws2 = wb.create_sheet("Note")
        ws2["A1"] = "Document confidentiel"
    out = analyze_workbook(build(f))
    assert out.tables == []
    assert sorted(codes(out)) == ["empty_sheet", "no_table_found", "no_table_found"]


def test_error_cells_are_reported():
    def f(wb):
        ws = wb.active
        ws.append(["Produit", "Prix", "Qté"])
        for i in range(12):
            ws.append([f"P{i}", 100 + i, i + 1])
        ws["B6"] = "#DIV/0!"
    out = analyze_workbook(build(f))
    fd = finding(out, "formula_errors")
    assert fd and fd.severity == "high" and fd.sample_rows == [6]


def test_negative_quantity_and_duplicate_ids():
    def f(wb):
        ws = wb.active
        ws.append(["id", "quantité", "prix"])
        for i in range(20):
            ws.append([f"A-{100 + (i if i != 7 else 6)}", 5 if i != 3 else -2, 10 + i])
    out = analyze_workbook(build(f))
    assert finding(out, "unexpected_negative_or_zero", "quantité").evidence["negatives"] == 1
    assert finding(out, "duplicate_ids", "id") is not None


def test_text_inside_numeric_column_is_invalid_not_silently_dropped():
    def f(wb):
        ws = wb.active
        ws.append(["Produit", "Prix"])
        for i in range(20):
            ws.append([f"P{i}", 100 + i])
        ws["B8"] = "gratuit"
    out = analyze_workbook(build(f))
    fd = finding(out, "invalid_values", "Prix")
    assert fd and fd.sample_rows == [8] and fd.evidence["examples"] == ["gratuit"]


def test_total_row_with_sum_formula_is_excluded():
    def f(wb):
        ws = wb.active
        ws.append(["Mois", "Ventes"])
        for i in range(6):
            ws.append([f"M{i}", 100 * (i + 1)])
        ws.append(["Total", "=SUM(B2:B7)"])
    out = analyze_workbook(build(f))
    t = out.tables[0]
    assert t.total_rows == [8] and len(t.rows) == 6


def test_limits_are_enforced():
    data = open(os.path.join(FIX, "01_sales_clean.xlsx"), "rb").read()
    with pytest.raises(Rejected) as e:
        analyze_workbook(data, Limits(max_cells=100))
    assert e.value.code == "too_many_cells"
    with pytest.raises(Rejected) as e:
        analyze_workbook(data, Limits(timeout_seconds=-1))
    assert e.value.code == "timeout"
    out = analyze_workbook(data, Limits(max_table_rows=50))
    assert out.tables[0].truncated and len(out.tables[0].rows) == 50 and "table_truncated" in codes(out)


def test_frames_match_the_profile():
    out = load("02_messy_ventes_fr.xlsx")
    df = out.frames[out.tables[0].id]
    assert len(df) == 61 and df["row_number"].iloc[0] == 5
    assert str(df["Date"].dtype).startswith("datetime64")
    assert df["Montant (DA)"].dtype == "float64"


# ------------------------------------------------------------------ formula helpers

def test_relativize_ignores_strings_and_functions():
    assert relativize("=A5*2", 5, 2) == "=R[0]C[-1]*2"
    assert relativize("=A6*2", 6, 2) == "=R[0]C[-1]*2"
    assert relativize('=IF(A5="A1",LOG10(B5),0)', 5, 3) == '=IF(R[0]C[-2]="A1",LOG10(R[0]C[-1]),0)'
    assert relativize("=$A$1*B5", 5, 3) == "=R1C1*R[0]C[-1]"


def test_inconsistent_formulas_found():
    f = {r: f"=A{r}*2" for r in range(2, 22)}
    f[10] = "=A10*3"
    res = inconsistent_formulas(f, 2)
    assert res["rows"] == [10] and res["pattern_share"] > 0.9
    assert inconsistent_formulas({r: f"=A{r}*2" for r in range(2, 22)}, 2) is None
    assert inconsistent_formulas({r: f"=A{r}*2" for r in range(2, 5)}, 2) is None
