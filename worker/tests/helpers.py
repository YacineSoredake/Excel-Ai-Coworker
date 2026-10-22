import io
from datetime import date, datetime

import openpyxl


def build(fn) -> bytes:
    wb = openpyxl.Workbook()
    fn(wb)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def col(out, table_idx=0):
    return {c.name: c for c in out.tables[table_idx].columns}


def codes(out, table_idx=None):
    return [f.code for f in out.findings if table_idx is None or f.table_index == table_idx]


def finding(out, code, column=None):
    for f in out.findings:
        if f.code == code and (column is None or (f.column_pos is not None and
                               out.tables[f.table_index].columns[f.column_pos].name == column)):
            return f
    return None
