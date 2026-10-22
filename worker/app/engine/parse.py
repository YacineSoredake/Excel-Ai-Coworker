"""Phase B — parse: workbook -> sheets, cells, formulas, merged cells, tables, names.

Two passes, as Excel needs: formulas (data_only=False) and cached values (data_only=True).
Macros are never executed and external links are never followed.
"""
import io
from dataclasses import dataclass, field

import openpyxl
from openpyxl.utils import get_column_letter, range_boundaries

from ..validate import Rejected
from .context import Ctx, Limits


@dataclass
class ExcelTable:
    name: str
    min_row: int  # 1-based
    min_col: int
    max_row: int
    max_col: int
    header_rows: int
    totals_rows: int


@dataclass
class SheetData:
    name: str
    position: int
    hidden: bool
    rows: list  # list[list[value]] — cached values, rectangular, 0-based
    n_rows: int
    n_cols: int
    formulas: dict  # (row, col) 1-based -> formula text
    errors: dict  # (row, col) 1-based -> '#DIV/0!' ...
    merged: list
    hyperlinks: int
    excel_tables: list
    _ws: object = field(default=None, repr=False)

    def number_format(self, row: int, col: int):
        try:
            return self._ws.cell(row=row, column=col).number_format
        except Exception:
            return None


@dataclass
class WorkbookData:
    sheets: list
    defined_names: list
    _wb: object = field(default=None, repr=False)

    def close(self):
        try:
            self._wb.close()
        except Exception:
            pass


def _guard_size(data: bytes, limits: Limits) -> None:
    ro = openpyxl.load_workbook(io.BytesIO(data), read_only=True)
    try:
        total = 0
        for ws in ro.worksheets:
            try:
                ws.calculate_dimension(force=True)
            except Exception:
                pass
            total += (ws.max_row or 0) * (ws.max_column or 0)
    finally:
        ro.close()
    if total > limits.max_cells:
        raise Rejected(
            "too_many_cells",
            f"This workbook is too large to analyze ({total:,} cells; the limit is {limits.max_cells:,}). "
            "Remove unused rows and columns, or split the file.",
        )


def _collect_formulas(data: bytes) -> dict:
    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=False, keep_links=False)
    out = {}
    try:
        for ws in wb.worksheets:
            f = {}
            for row in ws.iter_rows():
                for c in row:
                    if c.data_type == "f":
                        v = c.value
                        f[(c.row, c.column)] = v if isinstance(v, str) else str(getattr(v, "text", v))
            out[ws.title] = f
    finally:
        wb.close()
    return out


def _trim(rows: list) -> list:
    def empty(v):
        return v is None or (isinstance(v, str) and not v.strip())

    while rows and all(empty(v) for v in rows[-1]):
        rows.pop()
    if not rows:
        return rows
    width = len(rows[0])
    while width > 0 and all(empty(r[width - 1]) for r in rows):
        width -= 1
    return [r[:width] for r in rows]


def load_workbook_data(data: bytes, limits: Limits, ctx: Ctx) -> WorkbookData:
    _guard_size(data, limits)
    formulas = _collect_formulas(data)
    ctx.tick()
    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True, keep_links=False)
    sheets = []
    for i, ws in enumerate(wb.worksheets):
        ctx.tick()
        errors, rows = {}, []
        max_row, max_col = ws.max_row or 0, ws.max_column or 0
        if max_row and max_col:
            for row in ws.iter_rows(min_row=1, max_row=max_row, max_col=max_col):
                vals = []
                for c in row:
                    v = c.value
                    if c.data_type == "e":
                        errors[(c.row, c.column)] = str(v)
                        v = None
                    vals.append(v)
                rows.append(vals)
        rows = _trim(rows)
        tables = []
        try:
            for t in ws.tables.values():
                c0, r0, c1, r1 = range_boundaries(t.ref)
                tables.append(ExcelTable(
                    name=t.displayName or t.name, min_row=r0, min_col=c0, max_row=r1, max_col=c1,
                    header_rows=1 if (t.headerRowCount is None or t.headerRowCount >= 1) else 0,
                    totals_rows=int(t.totalsRowCount or 0)))
        except Exception:
            pass
        sheets.append(SheetData(
            name=ws.title, position=i, hidden=ws.sheet_state != "visible",
            rows=rows, n_rows=len(rows), n_cols=len(rows[0]) if rows else 0,
            formulas=formulas.get(ws.title, {}), errors=errors,
            merged=[r.coord for r in ws.merged_cells.ranges],
            hyperlinks=len(getattr(ws, "_hyperlinks", []) or []),
            excel_tables=tables, _ws=ws))
    names = []
    try:
        names = list(wb.defined_names.keys())[:50]
    except Exception:
        pass
    return WorkbookData(sheets=sheets, defined_names=names, _wb=wb)
