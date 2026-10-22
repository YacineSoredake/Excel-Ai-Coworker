"""Orchestrates phases B-D: parse -> detect tables -> normalize -> type -> roles -> profile -> findings."""
import time
import uuid
from collections import Counter

import pandas as pd
from openpyxl.utils import get_column_letter as L

from .columns import ColumnResult, analyze_column
from .context import Ctx, Limits
from .models import EngineOutput, Finding, TableResult
from .parse import SheetData, load_workbook_data
from .quality import table_findings
from .roles import assign_role
from .stats import basic_counts, detailed_stats
from .tables import TableCandidate, detect_tables, is_empty
from .textutil import clean_label, norm_key

PROFILE_VERSION = 1


def _dedupe(names: list[str]) -> list[str]:
    seen, out = Counter(), []
    for n in names:
        n = "row_number" if n == "__row" else n
        seen[n] += 1
        out.append(n if seen[n] == 1 else f"{n}_{seen[n]}")
    return out


def _label(v) -> str | None:
    if is_empty(v):
        return None
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return clean_label(v)


def build_table(sheet: SheetData, cand: TableCandidate, index: int, k_on_sheet: int, n_on_sheet: int,
                limits: Limits, ctx: Ctx) -> TableResult:
    rows_all = sheet.rows
    header = rows_all[cand.header_row] if cand.header_row is not None else None
    header_norm = [norm_key(header[c]) for c in range(cand.c0, cand.c1 + 1)] if header else []

    kept, blank, repeated = [], 0, 0
    for r in range(cand.data_r0, cand.data_r1 + 1):
        row = rows_all[r]
        if all(is_empty(row[c]) for c in range(cand.c0, cand.c1 + 1)):
            blank += 1
            continue
        if header and sum(1 for c, h in zip(range(cand.c0, cand.c1 + 1), header_norm)
                          if h and norm_key(row[c]) == h) >= max(2, 0.8 * len([h for h in header_norm if h])):
            repeated += 1
            continue
        kept.append(r)
    truncated = len(kept) > limits.max_table_rows
    kept = kept[: limits.max_table_rows]
    excel_rows = [r + 1 for r in kept]

    f_by_col: dict[int, set] = {}
    for (r, c) in sheet.formulas:
        f_by_col.setdefault(c, set()).add(r)

    labels = [_label(header[c]) if header else None for c in range(cand.c0, cand.c1 + 1)]
    cols_idx = list(range(cand.c0, cand.c1 + 1))
    keep_cols, dropped = [], 0
    for c, lab in zip(cols_idx, labels):
        if lab is None and all(is_empty(rows_all[r][c]) for r in kept):
            dropped += 1
            continue
        keep_cols.append((c, lab))
    names = _dedupe([lab if lab else f"column_{L(c + 1)}" for c, lab in keep_cols])

    columns: list[ColumnResult] = []
    for pos, ((c, lab), name) in enumerate(zip(keep_cols, names)):
        ctx.tick()
        raw = [rows_all[r][c] for r in kept]
        fmts = [sheet.number_format(r + 1, c + 1) for r, v in list(zip(kept, raw))[:400] if not is_empty(v)][:50]
        col = analyze_column(position=pos, letter=L(c + 1), name=name, original_label=lab, raw=raw,
                             rows=excel_rows, formats=fmts, formula_rows=f_by_col.get(c + 1, set()))
        basic_counts(col)
        assign_role(col)
        detailed_stats(col)
        columns.append(col)

    first_hdr = (cand.header_row + 1) if cand.header_row is not None else cand.data_r0 + 1
    last_row = (kept[-1] + 1) if kept else cand.data_r1 + 1
    name = cand.name or (sheet.name if n_on_sheet == 1 else f"{sheet.name} · table {k_on_sheet}")
    return TableResult(
        index=index, id=str(uuid.uuid4()), sheet=sheet.name, name=name, source=cand.source,
        cell_range=f"{L(cand.c0 + 1)}{first_hdr}:{L(cand.c1 + 1)}{last_row}",
        header_row=None if cand.header_row is None else cand.header_row + 1,
        first_data_row=cand.data_r0 + 1, last_data_row=last_row, col0=cand.c0 + 1, col1=cand.c1 + 1,
        confidence=cand.confidence, header_conf=cand.header_conf, columns=columns, rows=excel_rows,
        total_rows=[r + 1 for r in cand.total_rows], blank_rows=blank,
        repeated_header_rows=repeated, truncated=truncated, dropped_empty_columns=dropped,
        preamble=cand.preamble)


def to_frame(t: TableResult) -> pd.DataFrame:
    data = {"row_number": pd.Series(t.rows, dtype="int64")}
    for c in t.columns:
        vals = c.values
        if c.base_type in ("integer", "decimal", "currency", "percentage"):
            s = pd.Series([float("nan") if v is None else v for v in vals], dtype="float64")
        elif c.base_type in ("date", "datetime"):
            s = pd.to_datetime(pd.Series(vals, dtype="object"), errors="coerce")
        elif c.base_type == "boolean":
            s = pd.Series(vals, dtype="boolean")
        else:
            s = pd.Series([None if v is None else str(v) for v in vals], dtype="string")
        data[c.name] = s
    return pd.DataFrame(data)


def _keyset(c: ColumnResult) -> set:
    return {norm_key(v) for v in c.values if v is not None and str(v).strip()}


def detect_relationships(tables: list[TableResult]) -> list[dict]:
    """Probable foreign keys: most distinct values of A also appear in a key-like column B."""
    cands = []
    for t in tables:
        for c in t.columns:
            if c.role in ("dimension", "identifier") and c.base_type in ("text", "integer", "identifier") \
                    and c.distinct_count >= 2 and c.role_subtype not in ("flag", "ordinal", "period"):
                cands.append((t, c, _keyset(c)))
    out = []
    for tb, cb, kb in cands:
        if cb.uniqueness < 0.95 or len(kb) < 3:
            continue
        for ta, ca, ka in cands:
            if ta.id == tb.id or len(ka) < 2:
                continue
            contain = len(ka & kb) / len(ka)
            if contain >= 0.9:
                same_name = norm_key(ca.name) == norm_key(cb.name)
                out.append({"from_table": ta.id, "from_column": ca.name, "to_table": tb.id, "to_column": cb.name,
                            "from_table_name": ta.name, "to_table_name": tb.name,
                            "kind": "probable_foreign_key", "containment": round(contain, 3),
                            "confidence": round(min(0.99, 0.6 * contain + (0.3 if same_name else 0.1)), 2)})
    return sorted(out, key=lambda r: -r["confidence"])[:50]


def analyze_workbook(data: bytes, limits: Limits | None = None, has_external_links: bool = False) -> EngineOutput:
    limits = limits or Limits()
    ctx = Ctx(limits)
    t0 = time.monotonic()
    wbd = load_workbook_data(data, limits, ctx)
    sheets_out, tables, findings, frames = [], [], [], {}
    try:
        for sheet in wbd.sheets:
            ctx.tick()
            cands, notes = detect_tables(sheet)
            idxs = []
            for k, cand in enumerate(cands, start=1):
                t = build_table(sheet, cand, len(tables), k, len(cands), limits, ctx)
                tables.append(t)
                idxs.append(t.index)
                frames[t.id] = to_frame(t)
                findings.extend(table_findings(t, sheet))
            notes = (notes + [p for c in cands for p in c.preamble])[:8]
            cells = sum(1 for row in sheet.rows for v in row if not is_empty(v))
            sheets_out.append({
                "position": sheet.position, "name": sheet.name, "hidden": sheet.hidden,
                "row_count": sheet.n_rows, "column_count": sheet.n_cols, "cell_count": cells,
                "formula_count": len(sheet.formulas), "error_count": len(sheet.errors),
                "merged_count": len(sheet.merged), "hyperlink_count": sheet.hyperlinks,
                "is_empty": cells == 0, "notes": notes, "table_indexes": idxs})
            # sheet-level findings
            def sf(code, sev, title, detail, **kw):
                findings.append(Finding(code, sev, title, detail, sheet=sheet.name, **kw))
            if sheet.hidden:
                sf("hidden_sheet", "info", f"Sheet “{sheet.name}” is hidden",
                   "Hidden sheets are analyzed like any other sheet. Check that you want them included.")
            if cells == 0:
                sf("empty_sheet", "info", f"Sheet “{sheet.name}” is empty", "It contains no values.")
            elif not cands:
                sf("no_table_found", "info", f"No data table found on “{sheet.name}”",
                   "This sheet holds only a few notes or labels, so it was not analyzed as data.")
            in_table = lambda r, c: any(t.sheet == sheet.name and t.first_data_row - 1 <= r <= t.last_data_row
                                        and t.col0 <= c <= t.col1 for t in tables)
            stray = [(r, c) for (r, c) in sheet.errors if not in_table(r, c)]
            if stray:
                sf("formula_errors", "high", f"Potential issue: {len(stray)} cell(s) outside the data show an Excel error",
                   "These cells are not part of a detected table.", affected=len(stray),
                   sample_rows=sorted({r for r, _ in stray})[:10],
                   evidence={"by_error": dict(Counter(sheet.errors[k] for k in stray))})
    finally:
        wbd.close()

    if has_external_links:
        findings.append(Finding("external_links", "info", "The workbook links to other files",
                                "Those links are never followed. Values come from the results saved in this file."))

    rels = detect_relationships(tables)
    sev = Counter(f.severity for f in findings)
    main = max(tables, key=lambda t: len(t.rows), default=None)

    def cols(role):
        return [{"table": t.name, "column": c.name, "subtype": c.role_subtype}
                for t in tables for c in t.columns if c.role == role][:40]

    summary = {
        "profile_version": PROFILE_VERSION,
        "sheet_count": len(sheets_out), "table_count": len(tables), "total_rows": sum(len(t.rows) for t in tables),
        "main_table": {"id": main.id, "name": main.name} if main else None,
        "measures": cols("measure"), "dimensions": cols("dimension"), "identifiers": cols("identifier"),
        "time": [{"table": t.name, "column": c.name, "min": c.stats.get("min"), "max": c.stats.get("max"),
                  "granularity": c.stats.get("granularity")}
                 for t in tables for c in t.columns if c.role == "time"][:10],
        "findings_by_severity": {k: sev.get(k, 0) for k in ("high", "medium", "low", "info")},
        "defined_names": wbd.defined_names[:20],
    }
    return EngineOutput(sheets=sheets_out, tables=tables, findings=findings, relationships=rels,
                        summary=summary, frames=frames, engine_ms=int((time.monotonic() - t0) * 1000))
